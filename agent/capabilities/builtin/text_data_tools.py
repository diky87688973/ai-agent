# -*- coding: utf-8 -*-
"""文本与数据工具：text_diff / data_table。

- text_diff：两段文本 unified diff（左右各取文件或内联文本，只读不落盘）；
- data_table：csv/tsv/xlsx 的预览/筛选/排序/统计/导出（xlsx 依赖 openpyxl，可选降级）。
"""
from __future__ import annotations

import csv
import difflib
import io
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from agent.capabilities.registry import ToolPlugin


def _read_text_auto(fp: Path, encoding: str) -> str:
    """文件读取：显式编码 errors=replace；auto 按 utf-8→gb18030→gbk 尝试。"""
    if encoding != "auto":
        return fp.read_text(encoding=encoding, errors="replace")
    for enc in ("utf-8", "gb18030", "gbk"):
        try:
            return fp.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return fp.read_text(encoding="utf-8", errors="replace")


class TextDiffTool(ToolPlugin):
    """两段文本 unified diff。"""

    name = "text_diff"
    description = "对左右两侧文本生成 unified diff（只读不落盘）；左右各取「文件路径或内联文本」之一。"
    parameters = {
        "left_file": {"type": "string", "description": "左侧文件路径（与 left_text 二选一）"},
        "left_text": {"type": "string", "description": "左侧内联文本（与 left_file 二选一）"},
        "right_file": {"type": "string", "description": "右侧文件路径（与 right_text 二选一）"},
        "right_text": {"type": "string", "description": "右侧内联文本（与 right_file 二选一）"},
        "encoding": {"type": "string", "description": "读文件编码（默认 utf-8；auto 尝试链）"},
        "context": {"type": "integer", "description": "diff 上下文行数（默认 3，须 >= 0）"},
    }

    def _side_from_file_or_text(
        self,
        file_arg: Any,
        text_arg: Any,
        *,
        encoding: str,
    ) -> str:
        has_file = file_arg is not None and str(file_arg).strip() != ""
        has_text = text_arg is not None
        if int(has_file) + int(has_text) != 1:
            raise ValueError("每一侧必须且只能提供一个：文件路径或文本")
        if has_file:
            fp = Path(str(file_arg).strip())
            if not fp.is_file():
                raise FileNotFoundError(f"不是文件或不存在: {fp}")
            return _read_text_auto(fp, encoding)
        return str(text_arg)

    @staticmethod
    def _label_for_side(file_arg: Any, kind: str) -> str:
        if file_arg and str(file_arg).strip():
            try:
                return Path(str(file_arg).strip()).name
            except ValueError:
                return str(file_arg)
        return kind

    def execute(self, arguments: Dict[str, Any]) -> Dict[str, Any]:
        raw_context = arguments.get("context")
        try:
            context = int(raw_context) if raw_context is not None and str(raw_context).strip() != "" else 3
        except (TypeError, ValueError):
            raise ValueError("context 须为整数")
        if context < 0:
            raise ValueError("context 必须 >= 0")
        encoding = str(arguments.get("encoding") or "utf-8")
        left = self._side_from_file_or_text(arguments.get("left_file"), arguments.get("left_text"), encoding=encoding)
        right = self._side_from_file_or_text(arguments.get("right_file"), arguments.get("right_text"), encoding=encoding)
        left_lines = left.splitlines()
        right_lines = right.splitlines()
        from_label = self._label_for_side(arguments.get("left_file"), "left_text")
        to_label = self._label_for_side(arguments.get("right_file"), "right_text")
        diff_lines = list(
            difflib.unified_diff(
                left_lines,
                right_lines,
                fromfile=from_label,
                tofile=to_label,
                lineterm="",
                n=context,
            )
        )
        add_cnt = sum(1 for x in diff_lines if x.startswith("+") and not x.startswith("+++"))
        del_cnt = sum(1 for x in diff_lines if x.startswith("-") and not x.startswith("---"))
        summary = {
            "same": left == right,
            "left_chars": len(left),
            "right_chars": len(right),
            "added_lines": add_cnt,
            "deleted_lines": del_cnt,
        }
        diff_body = "\n".join(diff_lines)
        diff_md = "```diff\n" + diff_body + "\n```" if diff_lines else "```diff\n```"
        return {
            "summary": summary,
            "diff": diff_lines,
            "diff_markdown": diff_md,
        }


class DataTableTool(ToolPlugin):
    """表格读取与操作。"""

    name = "data_table"
    description = "表格文件（csv/tsv/xlsx）预览、筛选、排序、统计与导出（json/csv/markdown）；xlsx 需可选依赖 openpyxl。"
    parameters = {
        "action": {"type": "string", "description": "preview / filter / sort / stats / to_json / to_csv / to_md / sheets"},
        "path": {"type": "string", "description": "表格文件路径（.csv / .tsv / .xlsx）"},
        "sheet": {"type": "string", "description": "Excel 工作表名（xlsx 可选）"},
        "limit": {"type": "integer", "description": "预览行数上限（默认 20）"},
        "filter_col": {"type": "string", "description": "filter 的列名"},
        "filter_val": {"type": "string", "description": "filter 的匹配值"},
        "regex": {"type": "boolean", "description": "filter 时按正则匹配（默认等值忽略大小写）"},
        "sort_col": {"type": "string", "description": "sort 的列名"},
        "sort_desc": {"type": "boolean", "description": "sort 降序"},
        "col": {"type": "string", "description": "stats 的列名"},
    }

    @staticmethod
    def _load_csv(path: Path, *, delimiter: str = ",") -> List[Dict[str, Any]]:
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f, delimiter=delimiter)
            return list(reader)

    @staticmethod
    def _load_excel(path: Path, sheet: Optional[str]) -> List[Dict[str, Any]]:
        try:
            import openpyxl
        except ImportError:
            raise ImportError("读取 .xlsx 需要可选依赖 openpyxl：pip install openpyxl")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.active
            rows: List[Dict[str, Any]] = []
            headers: List[str] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(c) if c is not None else f"col{j}" for j, c in enumerate(row)]
                    continue
                row_dict: Dict[str, Any] = {}
                for j, val in enumerate(row):
                    col_name = headers[j] if j < len(headers) else f"col{j}"
                    row_dict[col_name] = val
                rows.append(row_dict)
            return rows
        finally:
            wb.close()

    @staticmethod
    def _detect_format(path: Path) -> str:
        s = path.suffix.lower()
        if s == ".xlsx":
            return "xlsx"
        if s == ".csv":
            return "csv"
        if s in (".tsv", ".tab"):
            return "tsv"
        return "unknown"

    def _load_table(self, path: Path, sheet: Optional[str]) -> Tuple[str, List[Dict[str, Any]]]:
        fmt = self._detect_format(path)
        if fmt == "xlsx":
            return fmt, self._load_excel(path, sheet)
        if fmt in ("csv", "tsv"):
            delim = "\t" if fmt == "tsv" else ","
            return fmt, self._load_csv(path, delimiter=delim)
        raise ValueError(f"不支持的文件格式: {path.suffix}（支持 .xlsx / .csv / .tsv）")

    @staticmethod
    def _do_preview(rows: List[Dict[str, Any]], limit: int) -> Dict[str, Any]:
        if not rows:
            return {"total_rows": 0, "columns": [], "preview": [], "limit": limit}
        columns = list(rows[0].keys())
        return {"total_rows": len(rows), "columns": columns, "preview": rows[:limit], "limit": limit}

    @staticmethod
    def _do_filter(rows: List[Dict[str, Any]], col: str, val: str, regex: bool) -> List[Dict[str, Any]]:
        if regex:
            pat = re.compile(val, re.IGNORECASE)
            return [r for r in rows if str(r.get(col, "")).strip() and pat.search(str(r.get(col, "")))]
        return [r for r in rows if str(r.get(col, "")).strip().lower() == val.lower()]

    @staticmethod
    def _do_sort(rows: List[Dict[str, Any]], col: str, desc: bool) -> List[Dict[str, Any]]:
        return sorted(
            rows,
            key=lambda r: str(r.get(col, "")) if r.get(col) is not None else "",
            reverse=desc,
        )

    @staticmethod
    def _do_stats(rows: List[Dict[str, Any]], col: str) -> Dict[str, Any]:
        values: List[float] = []
        for r in rows:
            v = r.get(col)
            if v is not None and v != "":
                try:
                    values.append(float(v))
                except (ValueError, TypeError):
                    pass
        stats: Dict[str, Any] = {"column": col, "total_rows": len(rows), "non_empty": len(values)}
        if values:
            stats["min"] = min(values)
            stats["max"] = max(values)
            stats["sum"] = round(sum(values), 6)
            stats["avg"] = round(sum(values) / len(values), 6)
        return stats

    def execute(self, arguments: Dict[str, Any]) -> Dict[str, Any]:
        act = str(arguments.get("action") or "").strip().lower().replace("-", "_")
        src_arg = str(arguments.get("path") or "").strip()
        if not src_arg:
            raise ValueError("缺少 path（表格文件路径）")
        src = Path(src_arg)
        if not src.is_file():
            raise FileNotFoundError(f"文件不存在: {src}")
        raw_limit = arguments.get("limit")
        try:
            limit = int(raw_limit) if raw_limit is not None and str(raw_limit).strip() != "" else 20
        except (TypeError, ValueError):
            raise ValueError("limit 须为整数")
        sheet = arguments.get("sheet")
        fmt, rows = self._load_table(src, str(sheet) if sheet else None)

        if act == "sheets":
            if fmt != "xlsx":
                return {"sheets": [src.name]}
            try:
                import openpyxl
            except ImportError:
                raise ImportError("读取 .xlsx 需要可选依赖 openpyxl：pip install openpyxl")
            wb = openpyxl.load_workbook(src, read_only=True)
            try:
                return {"sheets": wb.sheetnames}
            finally:
                wb.close()
        if act == "preview":
            return self._do_preview(rows, limit)
        if act == "filter":
            filter_col = arguments.get("filter_col")
            filter_val = arguments.get("filter_val")
            if not filter_col or filter_val is None:
                raise ValueError("filter 需要 filter_col 与 filter_val")
            filtered = self._do_filter(rows, str(filter_col), str(filter_val), bool(arguments.get("regex", False)))
            data = self._do_preview(filtered, limit)
            data["filtered_rows"] = len(filtered)
            return data
        if act == "sort":
            sort_col = arguments.get("sort_col")
            if not sort_col:
                raise ValueError("sort 需要 sort_col")
            return self._do_preview(self._do_sort(rows, str(sort_col), bool(arguments.get("sort_desc", False))), limit)
        if act == "stats":
            col = arguments.get("col")
            if not col:
                raise ValueError("stats 需要 col")
            return self._do_stats(rows, str(col))
        if act == "to_json":
            return {"content": json.dumps(rows, ensure_ascii=False, indent=2, default=str), "format": "json"}
        if act == "to_csv":
            if not rows:
                content = ""
            else:
                output = io.StringIO()
                writer = csv.DictWriter(output, fieldnames=list(rows[0].keys()))
                writer.writeheader()
                writer.writerows(rows)
                content = output.getvalue()
            return {"content": content, "format": "csv"}
        if act == "to_md":
            if not rows:
                content = "(空表)"
            else:
                cols = list(rows[0].keys())
                lines = ["|" + "|".join(cols) + "|", "|" + "|".join("---" for _ in cols) + "|"]
                for r in (rows[:limit] if limit else rows):
                    row_vals = []
                    for c in cols:
                        v = r.get(c)
                        row_vals.append("" if v is None else str(v).replace("|", "\\|"))
                    lines.append("|" + "|".join(row_vals) + "|")
                content = "\n".join(lines)
            return {"content": content, "format": "markdown"}
        raise ValueError(f"未知 action: {act}")


def build_plugins() -> List[ToolPlugin]:
    """构造本组工具（text_diff / data_table）。"""
    return [TextDiffTool(), DataTableTool()]
